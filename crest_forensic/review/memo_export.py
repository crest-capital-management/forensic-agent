"""Word export of a memo (Job 4). The status line sits in the page HEADER, so it is on every page.

Draft: "DRAFT - for human review". Signed: "Signed by <owner> on <date>". Nothing is sent anywhere; the file is only returned to the
person who asked for it.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _para(doc, text: str, bold: bool = False, size: int | None = None, style: str | None = None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return p


def build_docx(m: dict) -> bytes:
    c = m["content"]
    h = c["header"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    head = doc.sections[0].header.paragraphs[0]
    head.text = f"{m['banner']}   |   {m['reference']}   |   version {m['shown_version_no']}"
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in head.runs:
        r.bold = True
        r.font.size = Pt(9)

    _para(doc, "CREST Group", bold=True, size=11)
    doc.add_heading("Research assessment memo", level=0)
    _para(doc, f"{h['company']} ({h['symbol']})", bold=True, size=13)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List"
    for k, v in [("Reference", m["reference"]), ("Case", f"#{m['case_id']}"), ("Years in scope", ", ".join(h["years"])),
                 ("Basis", ", ".join(h["bases"])), ("Version", str(m["shown_version_no"])), ("Owner", h["owner"]),
                 ("Data as of", h["data_as_of"]), ("Status", m["banner"])]:
        row = table.add_row().cells
        row[0].text, row[1].text = k, v
    _para(doc, h["limits_line"], bold=not h["limits_approved"])
    for n in c.get("notices", []):
        _para(doc, "Notice: " + n)

    doc.add_heading("1. Assessment scope and objective", level=1)
    for t in c["scope"]:
        _para(doc, t)

    doc.add_heading("2. Background and trigger", level=1)
    for t in c["background"]:
        _para(doc, t)

    doc.add_heading("3. Findings summary", level=1)
    for i, f in enumerate(c["findings"], 1):
        doc.add_heading(f"Finding {i}: {f['check_id']}, {f['financial_year']}, {f['basis']} ({f['severity']}, {f['materiality']})", level=2)
        _para(doc, f["text"])
        _para(doc, f["limit_line"])
        r = f["reviewer"]
        _para(doc, f"Accepted by {r['by_name'] or r['by']} on {r['at'][:10]}: \"{r['reason']}\""
                   + (" (decision carried over from an earlier identical flag)" if r["carried_from_earlier_run"] else ""))
        _para(doc, "Sources: " + "; ".join(f"[{s['ref']}] {s['label']}" + (f" (quote: \"{s['quote']}\")" if s.get("quote") else "") for s in f["sources"]))
        _para(doc, f["standing_line"]).runs[0].italic = True

    doc.add_heading("4. Open questions and gaps", level=1)
    words = {"yes": "public data can close this", "partly": "public data can partly close this", "no": "public data cannot close this"}
    for q in c["open_questions"]:
        _para(doc, f"{q['text']} ({words[q['public_data']]}: {q['how']})", style="List Bullet")

    doc.add_heading("5. Recommended next steps", level=1)
    _para(doc, "Possible on public data", bold=True)
    for s in c["next_steps"]["public"]:
        _para(doc, s, style="List Bullet")
    _para(doc, "What a formal engagement with data access would need", bold=True)
    for s in c["next_steps"]["engagement"]:
        _para(doc, s, style="List Bullet")

    doc.add_heading("6. Ownership and sign-off", level=1)
    _para(doc, f"Owner: {h['owner']}")
    _para(doc, c["ownership"]["note"])
    _para(doc, (f"Signed by {m['owner_name']} on {m['signed_at'][:10]} (version {m['signed_version_no']})."
                if m["status"] == "signed" else "Signature: ______________________   Date: ______________"))

    doc.add_heading("Annex A. Method and limits", level=1)
    a = c["annex_a"]
    _para(doc, "Checks used: " + ", ".join(f"{x['check_id']} {x['name']}" for x in a["checks"]))
    for r in a["runs"]:
        _para(doc, f"Run {r['run_id']} ({r['basis']}), years {', '.join(r['years'])}, limits approved: {'yes' if r['config_approved'] else 'no'}, code {r['code_version']}", style="List Bullet")
    _para(doc, "Source precedence: " + a["source_rule"] + ".")
    _para(doc, a["library_line"])

    doc.add_heading("Annex B. Citations", level=1)
    seen = set()
    for s in c["annex_b"]:
        if s["ref"] in seen:
            continue
        seen.add(s["ref"])
        _para(doc, f"[{s['ref']}] {s['label']}", style="List Bullet")

    doc.add_heading("Annex C. Change log", level=1)
    for v in m["versions"]:
        _para(doc, f"Version {v['version_no']} ({v['kind']}), {v['created_at'][:16].replace('T', ' ')} UTC, by {v['created_by_name'] or v['created_by']}: {v['change_note']}", style="List Bullet")

    doc.add_heading("Annex D. Documentation references (ICAI FAIS, SEBI RIA)", level=1)
    _para(doc, c["annex_d"]["text"])

    e = c.get("annex_e")
    if e:
        doc.add_heading("Annex E. Scheme hypotheses (Job 3)", level=1)
        _para(doc, e["line"])
        for h in e["matches"]:
            _para(doc, f"{h['scheme_id']} {h['scheme_name']} ({h['class_label']}): hypothesis, {h['coverage_pct']}% of the scheme's indicators present. "
                       f"Accepted by {h['lead']}" + (f" and {h['partner']}" if h["partner"] else "") + ".", style="List Bullet")
            _para(doc, "Present: " + "; ".join(f"{p['indicator']} ({p['detail']})" for p in h["present"]))
            _para(doc, "Missing: " + ("; ".join(h["missing"]) or "none"))
            _para(doc, "Facts that cut against it: " + "; ".join(h["against"]))
            _para(doc, "Common innocent explanations: " + "; ".join(h["innocent"]))
            _para(doc, "Evidence that would confirm or refute it: " + "; ".join(f"{x['text']} ({'public' if x['public'] else 'not public'})" for x in h["evidence"]))
            _para(doc, "References: " + "; ".join(e["references"].get(h["scheme_id"], [])))

    for line in c["standing"]:
        _para(doc, line).runs[0].italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
