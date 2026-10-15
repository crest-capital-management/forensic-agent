import pytest
from sqlalchemy import select

from crest_forensic.db.models import AppUser, ConfigVersion, Flag, Run
from crest_forensic.review import memo as mm
from crest_forensic.review import memo_library as lib
from crest_forensic.review import service as svc
from crest_forensic.tests.test_review_api import client, company, session  # noqa: F401  (fixtures reused)


@pytest.fixture()
def users(session):
    svc.ensure_default_users(session)
    return {u.username: u for u in session.execute(select(AppUser)).scalars()}


def _inp(field, value):
    return {"name": field, "field": field, "fy": "FY2026", "value": value, "provider": "annual_report_pdf",
            "providers": ["annual_report_pdf"], "resolution": "AGREE", "snapshot_id": 1, "line_item_id": 1,
            "snapshot_fetched_at": "2026-10-04T10:43:33"}


@pytest.fixture()
def flag(session, company):
    cfg = ConfigVersion(version=1, content={}, content_hash="h")
    session.add(cfg)
    session.flush()
    run = Run(company_id=company.company_id, basis="standalone", years=["FY2026"], config_version_id=cfg.id,
              config_approved=False, code_version="t", run_signature="s", manifest={})
    session.add(run)
    session.flush()
    f = Flag(run_id=run.id, check_id="J1-03", check_version=1, financial_year="FY2026", severity="high", severity_rule_id="SEV-J1-03",
             what_observed="In FY2026, trade receivables changed +10.1% while revenue grew +2.8% (growth ratio 3.55x).",
             location={}, technique="t", materiality={"label": "material"}, evidence_confidence="high",
             figures={"thresholds_applied": {"high": {"value": 2.0, "source": "proposal"}, "medium": {"value": 1.5, "source": "proposal"}}},
             citation={"inputs": [_inp("trade_receivables", 5.6976e11), _inp("revenue_from_operations", 2.20938e12)]},
             content_hash="h1")
    session.add(f)
    session.flush()
    return f


def _case(session, users, company, portfolio=False, reason="routine screen"):
    return svc.create_case(session, users["analyst1"], company.company_id, ["FY2026"], reason, is_portfolio_company=portfolio)


def _accepted(session, users, company, flag, **kw):
    case = _case(session, users, company, **kw)
    svc.review_flag(session, users["analyst1"], case.id, flag.id, "accept", "figures agree with the filing")
    return case


def test_a_memo_needs_at_least_one_accepted_flag(session, users, company, flag):
    case = _case(session, users, company)
    with pytest.raises(svc.RuleError, match="accepted"):
        mm.create_memo(session, users["analyst1"], case.id, "owner1")
    svc.review_flag(session, users["analyst1"], case.id, flag.id, "reject", "explained in note 12")
    with pytest.raises(svc.RuleError, match="accepted"):
        mm.create_memo(session, users["analyst1"], case.id, "owner1")


def test_first_draft_has_the_sections_the_draft_line_citations_and_the_reviewer(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1", "revenue and receivables only")
    c = m["content"]
    assert m["status"] == "draft" and m["banner"] == mm.DRAFT_LINE and c["header"]["draft_line"] == mm.DRAFT_LINE
    assert m["reference"].startswith("CREST-TESTCO-") and m["reference"].endswith("-001") and c["header"]["version_no"] == 1
    assert "revenue and receivables only" in " ".join(c["scope"]) and "reaches no conclusion" in " ".join(c["scope"])
    (f,) = c["findings"]
    assert f["text"].startswith("In FY2026, trade receivables changed") and f["reviewer"]["by"] == "analyst1"
    assert f["reviewer"]["reason"] == "figures agree with the filing"
    refs = [s["ref"] for s in f["sources"]]
    assert refs[0] == f"flag:{flag.id}" and "fin:trade_receivables:FY2026:standalone" in refs and len(refs) == 3
    assert "draft" in f["limit_line"] and "high 2.0" in f["limit_line"]
    assert any("customers" in q["text"] for q in c["open_questions"])
    assert all(q["public_data"] in ("yes", "partly", "no") for q in c["open_questions"])
    assert c["next_steps"]["public"] and c["next_steps"]["engagement"]
    assert "has been approved" in c["annex_a"]["library_line"] and lib.LIBRARY_APPROVED
    assert "DRAFT" in c["header"]["limits_line"] and c["header"]["data_as_of"] == "2026-10-04"
    assert c["annex_d"]["supplied"] is True and "FAIS 120" in c["annex_d"]["text"] and "Regulation 25" in c["annex_d"]["text"] and c["header"]["owner"] == "Memo owner (dev)"


def test_unreviewed_and_rejected_flags_stay_out(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    other = Flag(run_id=flag.run_id, check_id="J1-04", check_version=1, financial_year="FY2026", severity="medium", severity_rule_id="R",
                 what_observed="other", location={}, technique="t", materiality={"label": "material"}, evidence_confidence="high",
                 figures={}, citation={"inputs": [_inp("inventories", 1e9)]}, content_hash="h2")
    session.add(other)
    session.flush()
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1")
    assert [f["flag_id"] for f in m["content"]["findings"]] == [flag.id]


def test_portfolio_company_flags_need_the_advisers_approval_first(session, users, company, flag):
    case = _accepted(session, users, company, flag, portfolio=True)
    with pytest.raises(svc.RuleError, match="adviser"):
        mm.create_memo(session, users["analyst1"], case.id, "owner1")
    esc = svc.request_escalation(session, users["analyst1"], case.id, flag.id, "needs a look")
    with pytest.raises(svc.RuleError):
        mm.create_memo(session, users["analyst1"], case.id, "owner1")        # escalated but not yet approved
    svc.approve_escalation(session, users["adviser1"], esc["id"])
    assert mm.create_memo(session, users["analyst1"], case.id, "owner1")["content"]["findings"]


def test_banned_wording_in_a_case_reason_or_a_review_reason_is_withheld_with_a_notice(session, users, company, flag):
    case = _case(session, users, company, reason="possible fraud at the company")
    svc.review_flag(session, users["analyst1"], case.id, flag.id, "accept", "this is a scam")
    c = mm.create_memo(session, users["analyst1"], case.id, "owner1")["content"]
    text = " ".join(mm._all_strings(c)).lower()
    assert "fraud" not in text and "scam" not in text and mm.WITHHELD in " ".join(c["scope"])
    assert len(c["notices"]) >= 2


def test_only_a_memo_owner_or_lead_can_be_named_owner(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    for bad in ("adviser1", "analyst1", "nobody"):
        with pytest.raises(svc.RuleError, match="owner"):
            mm.create_memo(session, users["analyst1"], case.id, bad)
    with pytest.raises(svc.RuleError):
        mm.create_memo(session, users["partner1"], case.id, "owner1")        # a partner may not draft


def test_an_edit_is_a_new_version_and_the_old_one_is_kept(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1")
    before = [dict(f) for f in m["content"]["findings"]]
    out = mm.edit_memo(session, users["owner1"], m["id"], {"background": ["A shorter background."],
                                                           "annex_d": "ICAI FAIS 100 and SEBI RIA record keeping, as supplied."}, "tightened background")
    assert out["latest_version_no"] == 2 and out["content"]["background"] == ["A shorter background."] and out["content"]["annex_d"]["supplied"]
    v1 = mm.get_memo(session, users["owner1"], m["id"], version_no=1)
    assert v1["content"]["background"] != ["A shorter background."] and v1["content"]["findings"] == before
    assert [(v["version_no"], v["kind"]) for v in out["versions"]] == [(1, "generated"), (2, "edited")]
    assert out["versions"][1]["change_note"] == "tightened background" and out["versions"][1]["created_by"] == "owner1"


def test_edits_are_checked_wording_numbers_notes_and_who(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1")
    fid = flag.id
    with pytest.raises(svc.RuleError, match="note"):
        mm.edit_memo(session, users["owner1"], m["id"], {"background": ["x"]}, " ")
    with pytest.raises(svc.RuleError, match="wording"):
        mm.edit_memo(session, users["owner1"], m["id"], {"background": ["This looks like fraud."]}, "n")
    with pytest.raises(svc.RuleError, match="number"):
        mm.edit_memo(session, users["owner1"], m["id"], {"findings": {fid: "Receivables grew 99% in FY2026."}}, "n")
    ok = mm.edit_memo(session, users["owner1"], m["id"], {"findings": {fid: "Trade receivables rose 10.1% in FY2026 while revenue rose 2.8%."}}, "simpler")
    assert ok["content"]["findings"][0]["text"].startswith("Trade receivables rose")
    with pytest.raises(svc.RuleError, match="cannot be edited"):
        mm.edit_memo(session, users["owner1"], m["id"], {"header": {"draft_line": ""}}, "n")      # the DRAFT marker is not editable
    with pytest.raises(svc.RuleError, match="Only the memo owner"):
        mm.edit_memo(session, users["analyst1"], m["id"], {"background": ["x"]}, "n")
    assert mm.edit_memo(session, users["lead1"], m["id"], {"background": ["lead edit"]}, "n")["latest_version_no"] == 3
    with pytest.raises(svc.RuleError, match="yes, partly or no"):
        mm.edit_memo(session, users["owner1"], m["id"], {"open_questions": [{"text": "q?", "public_data": "maybe"}]}, "n")


def test_a_changed_decision_makes_the_memo_out_of_date_and_blocks_signing_until_regenerated(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1")
    assert m["out_of_date"] is False
    svc.review_flag(session, users["analyst1"], case.id, flag.id, "reject", "second look: explained")
    assert mm.get_memo(session, users["owner1"], m["id"])["out_of_date"] is True
    with pytest.raises(svc.RuleError, match="changed"):
        mm.sign_memo(session, users["owner1"], m["id"])
    svc.review_flag(session, users["analyst1"], case.id, flag.id, "accept", "back to accepted")
    again = mm.regenerate(session, users["owner1"], m["id"], "decision changed")
    assert again["latest_version_no"] == 2 and again["out_of_date"] is False


def test_signing_is_by_the_named_owner_only_and_locks_the_memo(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1")
    for other in ("lead1", "analyst1"):
        with pytest.raises(svc.RuleError, match="named owner"):
            mm.sign_memo(session, users[other], m["id"])
    signed = mm.sign_memo(session, users["owner1"], m["id"])
    assert signed["status"] == "signed" and signed["banner"].startswith("Signed by Memo owner (dev) on ") and signed["signed_version_no"] == 1
    assert signed["can_edit"] is False and signed["can_sign"] is False
    for call in (lambda: mm.edit_memo(session, users["owner1"], m["id"], {"background": ["x"]}, "n"),
                 lambda: mm.regenerate(session, users["owner1"], m["id"]),
                 lambda: mm.sign_memo(session, users["owner1"], m["id"])):
        with pytest.raises(svc.RuleError, match="locked"):
            call()
    svc.review_flag(session, users["analyst1"], case.id, flag.id, "reject", "later")
    assert mm.get_memo(session, users["owner1"], m["id"])["out_of_date"] is False      # a signed memo is a record, not stale


def test_list_and_owner_choices(session, users, company, flag):
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "lead1")
    (row,) = mm.list_memos(session, case.id)
    assert row["id"] == m["id"] and row["status"] == "draft" and row["owner_name"] == "Forensic lead (dev)"
    assert {o["username"] for o in mm.owner_choices(session)} == {"owner1", "lead1"}


def test_a_finding_with_no_saved_source_is_refused(session, users, company, flag):
    flag.citation = {"inputs": [{"x": 1}]}
    session.flush()
    case = _accepted(session, users, company, flag)
    with pytest.raises(svc.RuleError, match="no saved figure"):
        mm.create_memo(session, users["analyst1"], case.id, "owner1")


# ---- over HTTP (the real routes, with the dev header) ---------------------------------------------------------------------

def _http_setup(client, session, company, flag):
    session.commit()
    h = {"X-User": "analyst1"}
    cid = client.post("/api/cases", headers=h, json={"company_id": company.company_id, "years": ["FY2026"], "reason": "screen"}).json()["id"]
    assert client.post(f"/api/cases/{cid}/flags/{flag.id}/review", headers=h, json={"decision": "accept", "reason": "agrees with filing"}).status_code == 200
    return cid, h


def test_memo_flow_over_http_with_roles_and_word_export(client, session, company, flag):
    cid, h = _http_setup(client, session, company, flag)
    assert client.post(f"/api/cases/{cid}/memos", headers={"X-User": "partner1"}, json={"owner": "owner1"}).status_code == 403
    assert client.post(f"/api/cases/{cid}/memos", headers=h, json={"owner": "adviser1"}).status_code == 400
    r = client.post(f"/api/cases/{cid}/memos", headers=h, json={"owner": "owner1", "scope_note": "receivables"})
    assert r.status_code == 200 and r.json()["status"] == "draft"
    mid = r.json()["id"]
    assert {o["username"] for o in client.get("/api/memo-owners", headers=h).json()} == {"lead1", "owner1"}
    assert client.get(f"/api/cases/{cid}/memos", headers=h).json()[0]["id"] == mid
    owner = {"X-User": "owner1"}
    assert client.put(f"/api/memos/{mid}", headers=h, json={"changes": {"background": ["x"]}, "note": "n"}).status_code == 403
    e = client.put(f"/api/memos/{mid}", headers=owner, json={"changes": {"background": ["Short background."]}, "note": "shorter"})
    assert e.status_code == 200 and e.json()["latest_version_no"] == 2
    assert client.get(f"/api/memos/{mid}?version=1", headers=h).json()["shown_version_no"] == 1
    word = client.get(f"/api/memos/{mid}/export.docx", headers=h)
    assert word.status_code == 200 and word.content[:2] == b"PK" and "wordprocessingml" in word.headers["content-type"]
    pdf = client.get(f"/api/memos/{mid}/export.pdf", headers=h)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF" and pdf.headers["content-type"] == "application/pdf"
    assert client.post(f"/api/memos/{mid}/sign", headers=h).status_code == 403
    signed = client.post(f"/api/memos/{mid}/sign", headers=owner)
    assert signed.status_code == 200 and signed.json()["status"] == "signed"
    assert client.put(f"/api/memos/{mid}", headers=owner, json={"changes": {"background": ["y"]}, "note": "n"}).status_code == 409


def test_the_word_file_carries_the_status_in_the_page_header_and_the_sections(session, users, company, flag):
    import io
    import docx
    from crest_forensic.review import memo_export
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1")
    d = docx.Document(io.BytesIO(memo_export.build_docx(m)))
    assert mm.DRAFT_LINE in d.sections[0].header.paragraphs[0].text and m["reference"] in d.sections[0].header.paragraphs[0].text
    text = "\n".join(p.text for p in d.paragraphs)
    for heading in ("1. Assessment scope", "2. Background", "3. Findings summary", "4. Open questions", "5. Recommended next steps",
                    "6. Ownership", "Annex A", "Annex B", "Annex C", "Annex D"):
        assert heading in text
    assert "trade receivables changed +10.1%" in text and "Statistical lead only" in text
    signed = mm.sign_memo(session, users["owner1"], m["id"])
    d2 = docx.Document(io.BytesIO(memo_export.build_docx(signed)))
    assert "Signed by Memo owner (dev)" in d2.sections[0].header.paragraphs[0].text


def test_the_pdf_has_the_status_line_on_every_page_and_all_the_sections(session, users, company, flag):
    import fitz
    from crest_forensic.review import memo_pdf
    case = _accepted(session, users, company, flag)
    m = mm.create_memo(session, users["analyst1"], case.id, "owner1")
    doc = fitz.open("pdf", memo_pdf.build_pdf(m))
    assert len(doc) >= 2
    for page in doc:
        assert mm.DRAFT_LINE in page.get_text() and m["reference"] in page.get_text()
    text = "\n".join(p.get_text() for p in doc)
    for heading in ("1. Assessment scope", "2. Background", "3. Findings summary", "4. Open questions", "5. Recommended next steps",
                    "6. Ownership", "Annex A", "Annex B", "Annex C", "Annex D"):
        assert heading in text
    assert "trade receivables changed +10.1%" in text.replace("\n", " ") and "Page 1 of" in text
    signed = mm.sign_memo(session, users["owner1"], m["id"])
    doc2 = fitz.open("pdf", memo_pdf.build_pdf(signed))
    assert all("Signed by Memo owner (dev)" in page.get_text() for page in doc2)


def test_the_memos_page_lists_every_memo_across_cases(client, session, company, flag):
    cid, h = _http_setup(client, session, company, flag)
    assert client.get("/api/memos", headers=h).json() == []
    mid = client.post(f"/api/cases/{cid}/memos", headers=h, json={"owner": "owner1"}).json()["id"]
    (row,) = client.get("/api/memos", headers=h).json()
    assert row["id"] == mid and row["company"] == "Test Co" and row["case_id"] == cid and row["status"] == "draft"
    assert row["reference"].startswith("CREST-TESTCO-") and row["versions"] == 1 and row["owner_name"] == "Memo owner (dev)"
